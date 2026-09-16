"""Bulk firewall inventory import (Settings > Asset Inventory).

Administrators download the *Firewall Inventory template* (.xlsx) and re-upload
it filled in. Every row carries an ``Action`` of ``add`` or ``remove``:

* ``add``    registers the firewall (device name + IP, with optional vendor,
  user name, password and port);
* ``remove`` removes the matching firewall from the inventory.

Each row is validated independently: an invalid or incomplete row is skipped
(reporting the offending column) while the remaining valid rows are applied.
"""

import io
import ipaddress

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from services import managed_firewalls_service

TEMPLATE_HEADERS = [
    "Serial Number",
    "Firewall Name",
    "Vendor",
    "IP Address",
    "User Name",
    "Password",
    "Port",
    "Action",
]

# Column aliases accepted on upload (normalised: lower-case, no spaces/-/_).
HEADER_ALIASES = {
    "serialnumber": "serial",
    "sno": "serial",
    "serial": "serial",
    "firewallname": "name",
    "devicename": "name",
    "name": "name",
    "vendor": "vendor",
    "ipaddress": "ip",
    "ip": "ip",
    "username": "username",
    "user": "username",
    "password": "password",
    "port": "port",
    "action": "action",
}

REQUIRED_HEADERS = ("name", "ip", "action")

MAX_ROWS = 2000


def build_template_bytes():
    """Return the Firewall Inventory template workbook as a ``BytesIO``."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Firewall Inventory"

    header_fill = PatternFill("solid", fgColor="FF1F2937")
    header_font = Font(bold=True, color="FFFFFFFF")

    for column, title in enumerate(TEMPLATE_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.fill = header_fill
        cell.font = header_font

    widths = [14, 22, 18, 18, 18, 18, 10, 12]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _normalise(value):
    text = str(value or "").strip().lower()
    for token in (" ", "_", "-", "\t", "\n"):
        text = text.replace(token, "")
    return text


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_indexes(header_row):
    indexes = {}
    for position, value in enumerate(header_row):
        key = HEADER_ALIASES.get(_normalise(value))
        if key and key not in indexes:
            indexes[key] = position
    return indexes


def _failure_message(name, column):
    """Build the user-facing failure message for a bad row."""

    return "Failure !! action cannot be perform on {0} as incorrect input for {1}.".format(
        name or "this row", column
    )


def _validate(row, indexes):
    """Return ``(fields, error_column)`` for one spreadsheet row.

    ``error_column`` is ``None`` when the row is valid.
    """

    def cell(key):
        position = indexes.get(key)
        if position is None or position >= len(row):
            return ""
        return _clean(row[position])

    fields = {
        "name": cell("name"),
        "ip": cell("ip"),
        "vendor": cell("vendor"),
        "username": cell("username"),
        "password": cell("password"),
        "port": cell("port"),
    }
    action = cell("action").lower()

    if not fields["name"]:
        return fields, "Firewall Name"
    if not action:
        return fields, "Action"
    if action not in ("add", "remove"):
        return fields, "Action"
    if not fields["ip"]:
        return fields, "IP Address"

    try:
        ipaddress.ip_address(fields["ip"])
    except ValueError:
        return fields, "IP Address"

    if fields["port"]:
        try:
            port = int(fields["port"])
        except (TypeError, ValueError):
            return fields, "Port"
        if not 1 <= port <= 65535:
            return fields, "Port"
        fields["port"] = port
    else:
        fields["port"] = None

    fields["action"] = action
    return fields, None


def process_upload(file_storage):
    """Validate and apply an uploaded Firewall Inventory workbook.

    Returns ``{"added", "removed", "failed", "results"}`` where each result is
    ``{"row", "name", "action", "status", "message"}``.
    """

    if file_storage is None or not (file_storage.filename or "").strip():
        raise ValueError("An Excel (.xlsx) file is required.")
    if not file_storage.filename.lower().endswith(".xlsx"):
        raise ValueError("Only Excel (.xlsx) files are supported.")

    try:
        workbook = load_workbook(file_storage.stream, read_only=True, data_only=True)
    except Exception:
        raise ValueError("The uploaded file could not be read as an Excel workbook.")

    sheet = workbook.active
    if sheet is None:
        raise ValueError("The uploaded workbook has no worksheet.")

    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    header_index = None
    indexes = {}
    for position, row in enumerate(rows):
        candidate = _header_indexes(row)
        if all(key in candidate for key in REQUIRED_HEADERS):
            header_index = position
            indexes = candidate
            break

    if header_index is None:
        raise ValueError(
            "The template must include the Firewall Name, IP Address and "
            "Action columns."
        )

    data_rows = rows[header_index + 1:]
    if not data_rows:
        raise ValueError("The uploaded file has no data rows.")

    results = []
    added = removed = failed = 0
    processed = 0

    for offset, row in enumerate(data_rows):
        if all(_clean(value) == "" for value in row):
            continue
        if processed >= MAX_ROWS:
            break
        processed += 1

        fields, error_column = _validate(row, indexes)

        if error_column:
            failed += 1
            results.append(
                {
                    "row": offset + header_index + 2,
                    "name": fields["name"],
                    "action": fields.get("action") or "",
                    "status": "error",
                    "message": _failure_message(fields["name"], error_column),
                }
            )
            continue

        try:
            if fields["action"] == "add":
                managed_firewalls_service.import_firewall(
                    fields["name"],
                    fields["ip"],
                    vendor=fields["vendor"],
                    port=fields["port"],
                    username=fields["username"],
                    password=fields["password"],
                )
                added += 1
                message = "Success !! {0} has been added.".format(fields["name"])
            else:
                managed_firewalls_service.remove_by_device_name(fields["name"])
                removed += 1
                message = "Success !! {0} has been removed.".format(fields["name"])
        except ValueError as exc:
            failed += 1
            results.append(
                {
                    "row": offset + header_index + 2,
                    "name": fields["name"],
                    "action": fields["action"],
                    "status": "error",
                    "message": _failure_message(fields["name"], "Firewall Name")
                    + " ({0})".format(str(exc)),
                }
            )
            continue

        results.append(
            {
                "row": offset + header_index + 2,
                "name": fields["name"],
                "action": fields["action"],
                "status": "success",
                "message": message,
            }
        )

    if processed == 0:
        raise ValueError("The uploaded file has no data rows.")

    return {
        "added": added,
        "removed": removed,
        "failed": failed,
        "results": results,
    }
